"""Génère la présentation PowerPoint sur l'API Slack."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- Palette ----------
BG = RGBColor(0x0C, 0x0F, 0x14)
PANEL = RGBColor(0x16, 0x1B, 0x24)
PANEL2 = RGBColor(0x1B, 0x21, 0x2C)
BORDER = RGBColor(0x26, 0x2E, 0x3A)
TEXT = RGBColor(0xEE, 0xF1, 0xF6)
TEXT_DIM = RGBColor(0x9A, 0xA5, 0xB4)
TEXT_FAINT = RGBColor(0x6B, 0x76, 0x86)
PURPLE = RGBColor(0x7C, 0x5C, 0xFF)
CYAN = RGBColor(0x36, 0xC5, 0xF0)
GREEN = RGBColor(0x2E, 0xB6, 0x7D)
YELLOW = RGBColor(0xEC, 0xB2, 0x2E)
PINK = RGBColor(0xE0, 0x1E, 0x5A)

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def set_no_autofit(tf):
    el = tf._txBody
    bodyPr = el.find(qn('a:bodyPr'))
    for tag in ('a:normAutofit', 'a:spAutoFit'):
        e = bodyPr.find(qn(tag))
        if e is not None:
            bodyPr.remove(e)
    bodyPr.append(bodyPr.makeelement(qn('a:noAutofit'), {}))


def box(slide, x, y, w, h, fill=None, line_color=None, line_w=Pt(1), dash=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape_type, x, y, w, h)
    if radius:
        try:
            sp.adjustments[0] = 0.06
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = line_w
        if dash:
            ln = sp.line._get_or_add_ln()
            d = ln.makeelement(qn('a:prstDash'), {'val': dash})
            ln.append(d)
    sp.shadow.inherit = False
    return sp


def text(slide, x, y, w, h, runs, size=14, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, font="Inter", line_spacing=1.15, wrap=True):
    """runs: string OR list of (text, dict-overrides)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    set_no_autofit(tf)
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    if isinstance(runs, str):
        runs = [(runs, {})]
    for txt, ov in runs:
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(ov.get('size', size))
        r.font.bold = ov.get('bold', bold)
        r.font.name = ov.get('font', font)
        r.font.color.rgb = ov.get('color', color)
    return tb


def multiline(slide, x, y, w, h, lines, size=13, color=TEXT_DIM, font="Inter",
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.25, space_after=6):
    """lines: list of (runs-list, paragraph-overrides-dict)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    set_no_autofit(tf)
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, (runs, pov) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = pov.get('align', align)
        p.line_spacing = pov.get('line_spacing', line_spacing)
        p.space_after = Pt(pov.get('space_after', space_after))
        if isinstance(runs, str):
            runs = [(runs, {})]
        for txt, ov in runs:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(ov.get('size', size))
            r.font.bold = ov.get('bold', False)
            r.font.name = ov.get('font', font)
            r.font.color.rgb = ov.get('color', color)
    return tb


def chip(slide, x, y, w, h, label, fill, text_color):
    sp = box(slide, x, y, w, h, fill=fill, radius=True)
    try:
        sp.adjustments[0] = 0.5
    except Exception:
        pass
    tf = sp.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = text_color
    r.font.name = "Inter"
    return sp


def kicker(slide, idx_label, title_label):
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), Inches(0.42), Inches(0.12), Inches(0.12))
    dot.fill.solid(); dot.fill.fore_color.rgb = PURPLE; dot.line.fill.background(); dot.shadow.inherit = False
    text(slide, Inches(0.78), Inches(0.34), Inches(8), Inches(0.3),
         [(idx_label + "  —  " + title_label, {})], size=12, bold=True, color=TEXT_FAINT)


def heading(slide, title, subtitle):
    text(slide, Inches(0.55), Inches(0.68), Inches(12), Inches(0.7),
         [(title, {})], size=30, bold=True, color=TEXT)
    text(slide, Inches(0.55), Inches(1.32), Inches(11.4), Inches(0.55),
         [(subtitle, {})], size=13, color=TEXT_DIM, line_spacing=1.3)


def panel(slide, x, y, w, h, title=None, icon_color=None, icon_bg=None, icon_char="●"):
    p = box(slide, x, y, w, h, fill=PANEL, line_color=BORDER, line_w=Pt(1), radius=True)
    cursor_y = y + Inches(0.22)
    if title:
        ic = box(slide, x + Inches(0.22), cursor_y, Inches(0.32), Inches(0.32),
                  fill=icon_bg or PANEL2, radius=True)
        try:
            ic.adjustments[0] = 0.3
        except Exception:
            pass
        itf = ic.text_frame
        itf.margin_left = 0; itf.margin_right = 0; itf.margin_top = 0; itf.margin_bottom = 0
        itf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ip = itf.paragraphs[0]; ip.alignment = PP_ALIGN.CENTER
        ir = ip.add_run(); ir.text = icon_char; ir.font.size = Pt(13); ir.font.color.rgb = icon_color or TEXT
        text(slide, x + Inches(0.64), cursor_y + Inches(0.01), w - Inches(0.9), Inches(0.32),
             [(title, {})], size=14, bold=True, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
        cursor_y += Inches(0.46)
    return cursor_y


def shot_placeholder(slide, x, y, w, h, label, sub, tag="Screenshot", icon="🖼"):
    sp = box(slide, x, y, w, h, fill=RGBColor(0x11, 0x15, 0x1C), line_color=BORDER, line_w=Pt(1.4),
              dash="dash", radius=True)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2); tf.margin_right = Inches(0.2)
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    r0 = p0.add_run(); r0.text = icon; r0.font.size = Pt(22)
    p1 = tf.add_paragraph(); p1.alignment = PP_ALIGN.CENTER; p1.space_before = Pt(6)
    r1 = p1.add_run(); r1.text = label; r1.font.size = Pt(12.5); r1.font.bold = True; r1.font.color.rgb = TEXT_DIM
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(3)
    r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(10); r2.font.color.rgb = TEXT_FAINT
    # corner tag
    tagbox = box(slide, x + w - Inches(1.5), y + Inches(0.1), Inches(1.35), Inches(0.28),
                 fill=RGBColor(0x00, 0x00, 0x00), radius=True)
    tagbox.fill.fore_color.rgb = RGBColor(0x00, 0x00, 0x00)
    tagbox.line.color.rgb = BORDER
    tagbox.line.width = Pt(0.75)
    ttf = tagbox.text_frame
    ttf.margin_left = 0; ttf.margin_right = 0; ttf.margin_top = 0; ttf.margin_bottom = 0
    ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tp = ttf.paragraphs[0]; tp.alignment = PP_ALIGN.CENTER
    tr = tp.add_run(); tr.text = tag.upper(); tr.font.size = Pt(8); tr.font.bold = True; tr.font.color.rgb = TEXT_FAINT
    return sp


def code_chip(slide, x, y, w, h, label, color, bg):
    return chip(slide, x, y, w, h, label, bg, color)


def msg_chip(slide, x, y, w, h, label, body):
    sp = box(slide, x, y, w, h, fill=RGBColor(0x18, 0x1B, 0x2A), line_color=PURPLE, line_w=Pt(1), radius=True)
    sp.line.color.rgb = PURPLE
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12); tf.margin_bottom = Inches(0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]
    r1 = p1.add_run(); r1.text = label.upper(); r1.font.size = Pt(9.5); r1.font.bold = True; r1.font.color.rgb = CYAN
    p2 = tf.add_paragraph(); p2.space_before = Pt(3)
    r2 = p2.add_run(); r2.text = body; r2.font.size = Pt(12); r2.font.color.rgb = TEXT
    return sp


def diagram(slide, x, y, w, h, nodes):
    """nodes: list of dicts {icon, label, sub, color, bg} and arrow labels between"""
    n = len(nodes)
    node_w = Inches(1.7)
    arrow_w = (w - node_w * n) / (n - 1) if n > 1 else Inches(0)
    cx = x
    for i, nd in enumerate(nodes):
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx + (node_w - Inches(0.55)) / 2, y, Inches(0.55), Inches(0.55))
        circ.fill.solid(); circ.fill.fore_color.rgb = nd['bg']
        circ.line.fill.background(); circ.shadow.inherit = False
        ctf = circ.text_frame
        ctf.margin_left = 0; ctf.margin_right = 0; ctf.margin_top = 0; ctf.margin_bottom = 0
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run(); cr.text = nd['icon']; cr.font.size = Pt(18)
        text(slide, cx, y + Inches(0.62), node_w, Inches(0.3), [(nd['label'], {})],
             size=11.5, bold=True, color=TEXT, align=PP_ALIGN.CENTER)
        text(slide, cx, y + Inches(0.92), node_w, Inches(0.5), [(nd['sub'], {})],
             size=9, color=TEXT_FAINT, align=PP_ALIGN.CENTER, line_spacing=1.1)
        if i < n - 1:
            ax = cx + node_w
            arrow_label = nd.get('arrow_to', '')
            text(slide, ax, y + Inches(0.05), arrow_w, Inches(0.3), [("➜", {})],
                 size=16, color=TEXT_FAINT, align=PP_ALIGN.CENTER)
            text(slide, ax, y + Inches(0.34), arrow_w, Inches(0.3), [(arrow_label, {})],
                 size=8.5, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
        cx += node_w + arrow_w


def bullets(slide, x, y, w, lines, size=12.5, gap=Pt(8), bullet_color=CYAN, color=TEXT_DIM):
    tb = slide.shapes.add_textbox(x, y, w, Inches(2))
    tf = tb.text_frame
    tf.word_wrap = True
    set_no_autofit(tf)
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, (bold_part, rest) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        p.line_spacing = 1.25
        r0 = p.add_run(); r0.text = "●  "; r0.font.size = Pt(size - 2); r0.font.color.rgb = bullet_color
        if bold_part:
            rb = p.add_run(); rb.text = bold_part; rb.font.size = Pt(size); rb.font.bold = True; rb.font.color.rgb = TEXT
        if rest:
            rr = p.add_run(); rr.text = rest; rr.font.size = Pt(size); rr.font.color.rgb = color
    return tb


def footer_brand(slide):
    text(slide, SW - Inches(3.3), Inches(0.18), Inches(3.1), Inches(0.3),
         [("SLACK API · GUIDE VISUEL", {})], size=9, bold=True, color=TEXT_FAINT, align=PP_ALIGN.RIGHT)


# ======================================================================
# SLIDE 1 — VUE D'ENSEMBLE
# ======================================================================
s = add_slide()
footer_brand(s)
kicker(s, "01 / 06", "Vue d'ensemble")
heading(s, "Une Slack App, deux mondes",
        "Tout se configure dans la console Slack API, et tout s'utilise ensuite dans le workspace Slack.")

colW = Inches(5.55)
colH = Inches(4.4)
leftX = Inches(0.55)
rightX = Inches(7.1)
topY = Inches(2.05)

panel(s, leftX, topY, colW, colH)
text(s, leftX + Inches(0.22), topY + Inches(0.55), colW - Inches(0.5), Inches(0.3),
     [("A", {'color': PURPLE, 'bold': True, 'size': 16}), ("  Console Slack API", {'color': TEXT, 'bold': True, 'size': 15})])
text(s, leftX + Inches(0.22), topY + Inches(0.92), colW - Inches(0.5), Inches(0.4),
     [("api.slack.com/apps — centre de configuration technique.", {})], size=11, color=TEXT_DIM)
tags_a = ["Création de l'app", "Scopes", "Tokens", "Events", "Secrets"]
tx = leftX + Inches(0.22)
ty = topY + Inches(1.4)
for t in tags_a:
    w = Inches(0.18 + 0.082 * len(t))
    chip(s, tx, ty, w, Inches(0.34), t, PANEL2, TEXT_DIM)
    tx += w + Inches(0.12)
shot_placeholder(s, leftX + Inches(0.22), topY + Inches(2.0), colW - Inches(0.44), Inches(2.15),
                  "Dashboard Slack API", "Vue d'ensemble de l'app (api.slack.com/apps)", icon="🖥")

panel(s, rightX, topY, colW, colH)
text(s, rightX + Inches(0.22), topY + Inches(0.55), colW - Inches(0.5), Inches(0.3),
     [("B", {'color': CYAN, 'bold': True, 'size': 16}), ("  Workspace Slack", {'color': TEXT, 'bold': True, 'size': 15})])
text(s, rightX + Inches(0.22), topY + Inches(0.92), colW - Inches(0.5), Inches(0.4),
     [("L'app elle-même : le bot vit et agit ici.", {})], size=11, color=TEXT_DIM)
tags_b = ["Installation", "Autorisation", "Utilisation du bot"]
tx = rightX + Inches(0.22)
ty = topY + Inches(1.4)
for t in tags_b:
    w = Inches(0.18 + 0.082 * len(t))
    chip(s, tx, ty, w, Inches(0.34), t, PANEL2, TEXT_DIM)
    tx += w + Inches(0.12)
shot_placeholder(s, rightX + Inches(0.22), topY + Inches(2.0), colW - Inches(0.44), Inches(2.15),
                  "Installation / Autorisation", "Install to Workspace ou écran d'autorisation OAuth", icon="📲")

# center arrow
text(s, Inches(6.25), topY + Inches(1.9), Inches(0.85), Inches(0.5), [("⇄", {})], size=26, color=CYAN, align=PP_ALIGN.CENTER)

msg_chip(s, Inches(0.55), Inches(6.65), Inches(12.2), Inches(0.6),
         "Message central", "Configuration dans la console API, usage dans Slack.")

# ======================================================================
# SLIDE 2 — OAUTH & PERMISSIONS
# ======================================================================
s = add_slide()
footer_brand(s)
kicker(s, "02 / 06", "OAuth & Permissions")
heading(s, "Le centre des droits et des tokens",
        "Console : Votre App → OAuth & Permissions. C'est ici que se définissent les droits du bot et les tokens.")

leftW = Inches(6.3)
rightW = Inches(5.85)
leftX = Inches(0.55)
rightX = Inches(7.05)
topY = Inches(2.05)
panelH = Inches(4.35)

panel(s, leftX, topY, leftW, panelH, "OAuth & Permissions", icon_color=PURPLE, icon_bg=RGBColor(0x22,0x1c,0x3a), icon_char="🔐")
shot_placeholder(s, leftX + Inches(0.22), topY + Inches(0.6), leftW - Inches(0.44), panelH - Inches(0.85),
                  "Page OAuth & Permissions", "Bot Token Scopes, User Token Scopes, Install to Workspace")

ry = topY
panel(s, rightX, ry, rightW, Inches(2.0), "Ce qu'on configure", icon_color=CYAN, icon_bg=RGBColor(0x12,0x2a,0x32), icon_char="⚙")
bullets(s, rightX + Inches(0.22), ry + Inches(0.78), rightW - Inches(0.44), [
    ("Bot Token Scopes ", "— droits du bot (ex: chat:write)"),
    ("User Token Scopes ", "— droits au nom d'un utilisateur"),
    ("Install to Workspace ", "— déclenche l'installation"),
], size=11.5)

ry2 = topY + Inches(2.25)
panel(s, rightX, ry2, rightW, Inches(2.1), "Ce qu'on récupère", icon_color=GREEN, icon_bg=RGBColor(0x10,0x2a,0x20), icon_char="🔑")
bullets(s, rightX + Inches(0.22), ry2 + Inches(0.78), rightW - Inches(0.44), [
    ("Bot User OAuth Token ", "— xoxb-..."),
    ("User Token ", "— xoxp-..."),
], size=11.5, bullet_color=GREEN)

msg_chip(s, Inches(0.55), Inches(6.65), Inches(6.0), Inches(0.6),
         "À retenir", "Les scopes définissent ce que le token a le droit de faire.")
msg_chip(s, Inches(6.75), Inches(6.65), Inches(6.0), Inches(0.6),
         "À retenir", "Le token xoxb- est le token principal du bot.")

# ======================================================================
# SLIDE 3 — EVENT SUBSCRIPTIONS
# ======================================================================
s = add_slide()
footer_brand(s)
kicker(s, "03 / 06", "Event Subscriptions")
heading(s, "Rendre le bot réactif",
        "Console : Votre App → Event Subscriptions. Slack y apprend quoi envoyer, et où l'envoyer.")

leftW = Inches(5.85)
rightW = Inches(6.3)
leftX = Inches(0.55)
rightX = Inches(6.55)
topY = Inches(2.05)

panel(s, leftX, topY, leftW, Inches(2.7), "Event Subscriptions", icon_color=YELLOW, icon_bg=RGBColor(0x2b,0x24,0x10), icon_char="⚡")
shot_placeholder(s, leftX + Inches(0.22), topY + Inches(0.6), leftW - Inches(0.44), Inches(1.95),
                  "Page Event Subscriptions", "Enable Events, Request URL, Subscribe to Bot Events")

panel(s, rightX, topY, rightW, Inches(1.3), "Ce qu'on configure", icon_color=PURPLE, icon_bg=RGBColor(0x22,0x1c,0x3a), icon_char="⚙")
bullets(s, rightX + Inches(0.22), topY + Inches(0.78), rightW - Inches(0.44), [
    ("Enable Events ", "— active la fonctionnalité"),
    ("Request URL ", "— endpoint recevant les events"),
    ("Subscribe to Bot Events ", "— events écoutés"),
], size=10.5)

panel(s, rightX, topY + Inches(1.45), rightW, Inches(1.25), "Exemples d'events", icon_color=CYAN, icon_bg=RGBColor(0x12,0x2a,0x32), icon_char="📋")
tx = rightX + Inches(0.22)
ty = topY + Inches(1.45) + Inches(0.78)
for t in ["app_mention", "message.im", "message.channels"]:
    w = Inches(0.25 + 0.085 * len(t))
    chip(s, tx, ty, w, Inches(0.36), t, PANEL2, CYAN)
    tx += w + Inches(0.14)

panel(s, Inches(0.55), Inches(5.0), Inches(12.2), Inches(1.5), "Schéma du flux", icon_color=GREEN, icon_bg=RGBColor(0x10,0x2a,0x20), icon_char="🧭")
diagram(s, Inches(1.6), Inches(5.7), Inches(10.0), Inches(1.3), [
    {'icon': '💬', 'label': 'Slack', 'sub': 'Événement déclenché', 'bg': PURPLE, 'arrow_to': 'HTTP POST'},
    {'icon': '🌐', 'label': 'Request URL', 'sub': 'Endpoint configuré', 'bg': YELLOW, 'arrow_to': 'Reçu par'},
    {'icon': '🔁', 'label': 'Backend / n8n', 'sub': 'Traitement du workflow', 'bg': CYAN},
])

msg_chip(s, Inches(0.55), Inches(6.75), Inches(6.0), Inches(0.6), "À retenir", "Ne donne PAS les droits.")
msg_chip(s, Inches(6.75), Inches(6.75), Inches(6.0), Inches(0.6), "À retenir", "Dit ce que Slack ENVOIE à l'app.")

# ======================================================================
# SLIDE 4 — EVENTS + SCOPES
# ======================================================================
s = add_slide()
footer_brand(s)
kicker(s, "04 / 06", "Alignement")
heading(s, "Events + Scopes = comportement réel",
        "Le bot ne fonctionne que si les events souscrits et les scopes accordés sont alignés.")

tableX, tableY = Inches(0.55), Inches(2.05)
tableW, tableH = Inches(12.2), Inches(3.4)
rows, cols = 4, 3
gfx = s.shapes.add_table(rows, cols, tableX, tableY, tableW, tableH)
tbl = gfx.table
tbl.columns[0].width = Inches(4.4)
tbl.columns[1].width = Inches(2.7)
tbl.columns[2].width = Inches(5.1)

headers = ["Comportement souhaité", "Event à souscrire", "Scope(s) nécessaire(s)"]
data = [
    ("🔔 Réagir à une mention", "le bot répond quand on le tag", "app_mention", "app_mentions:read  +  chat:write"),
    ("✉ Répondre en message privé", "conversation en DM avec le bot", "message.im", "im:history, chat:write"),
    ("📖 Lire les messages d'un canal", "écoute passive d'un channel", "message.channels", "channels:history"),
]

for c, h in enumerate(headers):
    cell = tbl.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = BG
    cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.08); cell.margin_bottom = Inches(0.08)
    tf = cell.text_frame
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = h.upper()
    r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = TEXT_FAINT

for ridx, (title, sub, ev, sc) in enumerate(data, start=1):
    c0 = tbl.cell(ridx, 0)
    c0.fill.solid(); c0.fill.fore_color.rgb = PANEL2
    c0.margin_left = Inches(0.15); c0.margin_top = Inches(0.1)
    tf0 = c0.text_frame; tf0.word_wrap = True
    p0 = tf0.paragraphs[0]
    r0 = p0.add_run(); r0.text = title; r0.font.size = Pt(12.5); r0.font.bold = True; r0.font.color.rgb = TEXT
    p0b = tf0.add_paragraph(); p0b.space_before = Pt(2)
    r0b = p0b.add_run(); r0b.text = sub; r0b.font.size = Pt(9.5); r0b.font.color.rgb = TEXT_FAINT

    c1 = tbl.cell(ridx, 1)
    c1.fill.solid(); c1.fill.fore_color.rgb = PANEL2
    c1.margin_left = Inches(0.15); c1.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf1 = c1.text_frame
    p1 = tf1.paragraphs[0]
    r1 = p1.add_run(); r1.text = ev; r1.font.size = Pt(11.5); r1.font.bold = True; r1.font.color.rgb = CYAN; r1.font.name = "Consolas"

    c2 = tbl.cell(ridx, 2)
    c2.fill.solid(); c2.fill.fore_color.rgb = PANEL2
    c2.margin_left = Inches(0.15); c2.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf2 = c2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run(); r2.text = sc; r2.font.size = Pt(11); r2.font.color.rgb = TEXT_DIM; r2.font.name = "Consolas"

# remove default table style banding by setting all borders subtle (skip — acceptable default)

msg_chip(s, Inches(0.55), Inches(5.75), Inches(6.0), Inches(0.6), "Scopes", "= droits accordés au token")
msg_chip(s, Inches(6.75), Inches(5.75), Inches(6.0), Inches(0.6), "Events", "= déclencheurs entrants envoyés par Slack")
shot_placeholder(s, Inches(0.55), Inches(6.55), Inches(12.2), Inches(0.7),
                  "Liste d'events configurés (optionnel)", "Capture de la liste des Bot Events souscrits", icon="🗂")

# ======================================================================
# SLIDE 5 — SIGNING SECRET
# ======================================================================
s = add_slide()
footer_brand(s)
kicker(s, "05 / 06", "Signing Secret")
heading(s, "Vérifier que Slack est bien Slack",
        "Console : Basic Information → App Credentials → Signing Secret. Ce n'est pas un token d'accès.")

leftW = Inches(5.85)
rightW = Inches(6.3)
leftX = Inches(0.55)
rightX = Inches(6.55)
topY = Inches(2.05)

panel(s, leftX, topY, leftW, Inches(2.7), "Basic Information", icon_color=PINK, icon_bg=RGBColor(0x32,0x16,0x22), icon_char="🛡")
shot_placeholder(s, leftX + Inches(0.22), topY + Inches(0.6), leftW - Inches(0.44), Inches(1.95),
                  "App Credentials", "Section App Credentials avec le Signing Secret")

panel(s, rightX, topY, rightW, Inches(1.3), "Où le trouver", icon_color=PINK, icon_bg=RGBColor(0x32,0x16,0x22), icon_char="📍")
bullets(s, rightX + Inches(0.22), topY + Inches(0.78), rightW - Inches(0.44), [
    ("", "Basic Information > App Credentials > Signing Secret"),
], size=11.5, bullet_color=PINK)

panel(s, rightX, topY + Inches(1.45), rightW, Inches(1.25), "À quoi il sert", icon_color=PINK, icon_bg=RGBColor(0x32,0x16,0x22), icon_char="🎯")
bullets(s, rightX + Inches(0.22), topY + Inches(1.45) + Inches(0.55), rightW - Inches(0.44), [
    ("", "Vérifier que la requête vient bien de Slack"),
    ("", "Sécuriser les appels entrants vers la Request URL"),
], size=10.5, bullet_color=PINK)

panel(s, Inches(0.55), Inches(5.0), Inches(12.2), Inches(1.5), "Mini schéma de vérification", icon_color=CYAN, icon_bg=RGBColor(0x12,0x2a,0x32), icon_char="🧭")
diagram(s, Inches(1.6), Inches(5.7), Inches(10.0), Inches(1.3), [
    {'icon': '💬', 'label': 'Slack', 'sub': 'Envoie une requête signée', 'bg': PURPLE, 'arrow_to': 'Signature HMAC'},
    {'icon': '🔍', 'label': 'Vérification', 'sub': 'Avec le Signing Secret', 'bg': PINK, 'arrow_to': 'Si valide'},
    {'icon': '✅', 'label': 'Requête acceptée', 'sub': 'Traitement autorisé', 'bg': GREEN},
])

msg_chip(s, Inches(0.55), Inches(6.75), Inches(6.0), Inches(0.6), "Token", "= appeler Slack")
msg_chip(s, Inches(6.75), Inches(6.75), Inches(6.0), Inches(0.6), "Signing Secret", "= vérifier Slack quand Slack appelle ton système")

# ======================================================================
# SLIDE 6 — MAPPING N8N
# ======================================================================
s = add_slide()
footer_brand(s)
kicker(s, "06 / 06", "Mapping Slack ↔ n8n")
heading(s, "Quoi copier, et où",
        "Le pont concret entre la console Slack API et la configuration n8n.")

tableX, tableY = Inches(0.55), Inches(2.0)
tableW, tableH = Inches(12.2), Inches(2.2)
gfx2 = s.shapes.add_table(4, 3, tableX, tableY, tableW, tableH)
tbl2 = gfx2.table
tbl2.columns[0].width = Inches(4.0)
tbl2.columns[1].width = Inches(4.2)
tbl2.columns[2].width = Inches(4.0)
headers2 = ["Depuis Slack API", "Vers n8n", "Remarque"]
data2 = [
    ("🔑 Bot User OAuth Token", "OAuth & Permissions", "Champ Access Token du credential Slack", "Suffit souvent pour ENVOYER des messages"),
    ("🛡 Signing Secret", "Basic Information", "Champ Signature Secret du Slack Trigger", "Nécessaire si trigger / webhook entrant"),
    ("🌐 Request URL", "Event Subscriptions", "URL générée par le Slack Trigger n8n", "À coller dans Event Subscriptions"),
]
for c, h in enumerate(headers2):
    cell = tbl2.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = BG
    cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.06); cell.margin_bottom = Inches(0.06)
    tf = cell.text_frame
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = h.upper()
    r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = TEXT_FAINT

for ridx, (title, sub, mid, note) in enumerate(data2, start=1):
    c0 = tbl2.cell(ridx, 0)
    c0.fill.solid(); c0.fill.fore_color.rgb = PANEL2
    c0.margin_left = Inches(0.15); c0.margin_top = Inches(0.08)
    tf0 = c0.text_frame; tf0.word_wrap = True
    p0 = tf0.paragraphs[0]
    r0 = p0.add_run(); r0.text = title; r0.font.size = Pt(11.5); r0.font.bold = True; r0.font.color.rgb = TEXT
    p0b = tf0.add_paragraph(); p0b.space_before = Pt(2)
    r0b = p0b.add_run(); r0b.text = sub; r0b.font.size = Pt(9); r0b.font.color.rgb = TEXT_FAINT

    c1 = tbl2.cell(ridx, 1)
    c1.fill.solid(); c1.fill.fore_color.rgb = PANEL2
    c1.margin_left = Inches(0.15); c1.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf1 = c1.text_frame; tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    r1 = p1.add_run(); r1.text = mid; r1.font.size = Pt(10.5); r1.font.color.rgb = CYAN

    c2 = tbl2.cell(ridx, 2)
    c2.fill.solid(); c2.fill.fore_color.rgb = PANEL2
    c2.margin_left = Inches(0.15); c2.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf2 = c2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run(); r2.text = note; r2.font.size = Pt(10); r2.font.color.rgb = TEXT_DIM

panel(s, Inches(0.55), Inches(4.45), Inches(6.0), Inches(1.55), "Schéma global", icon_color=PURPLE, icon_bg=RGBColor(0x22,0x1c,0x3a), icon_char="🧭")
diagram(s, Inches(0.9), Inches(5.1), Inches(5.3), Inches(0.85), [
    {'icon': '🧩', 'label': 'Slack API', 'sub': 'Tokens, scopes, secret', 'bg': PURPLE, 'arrow_to': 'copie'},
    {'icon': '🗝', 'label': 'n8n Creds', 'sub': 'Slack API credential', 'bg': CYAN, 'arrow_to': 'utilise'},
])
text(s, Inches(0.9), Inches(5.95), Inches(5.0), Inches(0.3), [], size=1)

panel(s, Inches(6.75), Inches(4.45), Inches(6.0), Inches(1.55))
shot_placeholder(s, Inches(6.97), Inches(4.67), Inches(2.7), Inches(1.1), "n8n Credentials", "Slack API credential", icon="🗝")
shot_placeholder(s, Inches(9.85), Inches(4.67), Inches(2.7), Inches(1.1), "Slack Trigger", "Event Subscriptions côté n8n", icon="⚡")

msg_chip(s, Inches(0.55), Inches(6.2), Inches(6.0), Inches(0.95),
         "Envoyer depuis n8n", "Token + scopes suffisent souvent.")
msg_chip(s, Inches(6.75), Inches(6.2), Inches(6.0), Inches(0.95),
         "Déclencher depuis Slack", "Il faut aussi Event Subscriptions + idéalement Signing Secret.")

prs.save("/home/user/presentation_slack_api/Slack_API_Presentation.pptx")
print("OK")
